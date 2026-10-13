"""Unit tests for value parsers and normalization onto the import schema."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.services import excel, normalize, parsers


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("15/07/2026 (Jour 1) Matin 07h30-09h", "15/07/2026"),
        ("15 JUILLET 2026 : MATIN ET SOIR", "15/07/2026"),
        ("1er août 2026", "01/08/2026"),
        ("7h30-9h", ""),
    ],
)
def test_parse_date(text: str, expected: str) -> None:
    assert parsers.parse_date(text) == expected


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("15/07/2026 (Jour 1) Matin 07h30-09h", "07h30-09h00"),
        ("7h30-9h", "07h30-09h00"),
        ("14h-15h30", "14h00-15h30"),
        ("15/07/2026 SOIR", "SOIR"),
        ("17/07/2026 SOTR", "SOIR"),  # OCR misreading of SOIR
        ("15 JUILLET 2026 : MATIN ET SOIR", ""),
    ],
)
def test_parse_horaire(text: str, expected: str) -> None:
    assert parsers.parse_horaire(text) == expected


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("Jury N°1", "1"),
        ("Jury N°12", "12"),
        ("YAMWAYA /JURY III", "3"),
        ("BIMBILIN /JURY Il", "2"),  # II misread
        ("LPNK /JURY |", "1"),  # I misread
        ("II", "2"),
        ("IT", "2"),
        ("IV", "4"),
        ("YAMWAYA /JURY", ""),
        ("", ""),
    ],
)
def test_parse_jury(text: str, expected: str) -> None:
    assert parsers.parse_jury(text) == expected


def test_parse_lieu_drops_the_jury_part() -> None:
    assert parsers.parse_lieu("YAMWAYA /JURY II") == "YAMWAYA"
    assert parsers.parse_lieu("LYCEE SAINT- GABRIEL") == "LYCEE SAINT-GABRIEL"


def test_parse_tranche() -> None:
    assert parsers.parse_tranche("De Récépissé N°000023 à N°029985") == (
        "000023-029985",
        "000023",
        "029985",
    )
    assert parsers.parse_tranche("RECEPISSE N°061809A N°072958")[1:] == ("061809", "072958")
    assert parsers.parse_tranche("Toutes les candidates concernées")[0] == parsers.ALL_CANDIDATES
    assert parsers.parse_tranche("ous les candidats concernés")[0] == parsers.ALL_CANDIDATES
    assert parsers.parse_tranche("???") == ("", "", "")


def test_three_layouts_normalize_to_the_same_schema() -> None:
    """Combined date/hour column, banner-row date and split columns all map."""
    kaya = [
        ["Date et Horaires de Composition", "N°", "Intitulés du Concours", "Nombre de candidats par Concours",
         "Sous-centre (lieu) de Composition", "N° Jury", "Tranche de numéro récépissé de Candidats"],
        ["15/07/2026 (Jour 1) Matin 07h30-09h", "6", "Attachés d'Éducation", "1413",
         "Lycée Municipal de Kaya (LMK)", "Jury N°2", "De Récépissé N°029990 à N°055012"],
        ["15/07/2026 (Jour 1) Matin 07h30-09h", "Sous Total du 15/07/2026", "", "3544", "", "6", ""],
    ]  # fmt: skip
    ouahigouya = [
        ["N°", "HORAIRE", "INTITULE DES CONCOURS", "TRANCHES DES CANDIDATS PAR NUMEROS RECEPISSES",
         "JURY/LIEU", "NOMBRE DE SALLES"],
        ["16 JUILLET 2026 : MATIN", "", "", "", "", ""],
        ["27", "7h30-9h", "Maîtres des Activités Physiques", "Toutes les candidates concernées",
         "YAMWAYA /JURY I", "3"],
    ]  # fmt: skip
    dedougou = [
        ["DATE", "Intitulé du concours", "Nombre de Jury", "Lieu de composition", "Tranche de noms"],
        ["15/07/2026 SOIR", "Professeurs des Ecoles", "II", "LYCEE MUNICIPAL", "RECEPISSE N°000050 A N°026542"],
    ]  # fmt: skip

    rows = [normalize.normalize_pages([grid]).rows for grid in (kaya, ouahigouya, dedougou)]

    assert [len(r) for r in rows] == [1, 1, 1]  # header, banner and sub-total rows dropped
    assert rows[0][0].values == {
        "date": "15/07/2026", "horaire": "07h30-09h00", "concours": "Attachés d'Éducation",
        "lieu": "Lycée Municipal de Kaya (LMK)", "jury": "2",
        "tranche": "029990-055012", "tranche_debut": "029990", "tranche_fin": "055012",
    }  # fmt: skip
    assert rows[1][0].values["date"] == "16/07/2026"  # taken from the banner row
    assert (rows[1][0].values["lieu"], rows[1][0].values["jury"]) == ("YAMWAYA", "1")
    assert rows[1][0].values["tranche"] == parsers.ALL_CANDIDATES
    assert (rows[2][0].values["horaire"], rows[2][0].values["jury"]) == ("SOIR", "2")
    assert all(not row[0].doubtful for row in rows)


def test_page_without_header_continues_the_previous_table() -> None:
    header = ["DATE", "Intitulé du concours", "Nombre de Jury", "Lieu de composition", "Tranche de noms"]
    page1 = [header, ["15/07/2026 MATIN", "Conseillers d'Education", "II", "LYCEE MUNICIPAL", "Tous les candidats"]]
    page2 = [["15/07/2026 MATIN", "Attachés d'Education", "I", "HANKUY", "Tous les candidats"]]

    table = normalize.normalize_pages([page1, page2])

    assert [r.values["concours"] for r in table.rows] == ["Conseillers d'Education", "Attachés d'Education"]
    assert [r.page for r in table.rows] == [1, 2]
    assert table.warnings == []


def test_unrecognised_page_is_reported_not_guessed() -> None:
    table = normalize.normalize_pages([[["Produit", "Qte"], ["Pommes", "12"]]])
    assert table.rows == []
    assert len(table.warnings) == 1


def test_unparsed_fields_are_flagged_doubtful() -> None:
    grid = [
        ["DATE", "Intitulé du concours", "Nombre de Jury", "Lieu de composition", "Tranche de noms"],
        ["15/07/2026 MATIN", "Attachés d'Education", "??", "HANKUY", "illisible"],
    ]
    row = normalize.normalize_pages([grid]).rows[0]
    assert row.doubtful == ["jury", "tranche"]
    assert row.values["tranche"] == "illisible"  # raw text kept for manual fixing


def test_xlsx_opens_on_import_sheet_with_doubtful_cells_highlighted(tmp_path: Path) -> None:
    grid = [
        ["DATE", "Intitulé du concours", "Nombre de Jury", "Lieu de composition", "Tranche de noms"],
        ["15/07/2026 MATIN", "Attachés d'Education", "??", "HANKUY", "Tous les candidats"],
    ]
    table = normalize.normalize_pages([grid])
    out = tmp_path / "out.xlsx"
    excel.export_xlsx([grid], out, normalized=table)

    workbook = openpyxl.load_workbook(out)
    assert workbook.sheetnames == ["Import", "Page 1"]
    sheet = workbook["Import"]
    assert [c.value for c in sheet[1]] == table.columns
    jury = sheet.cell(row=2, column=table.columns.index("jury") + 1)
    lieu = sheet.cell(row=2, column=table.columns.index("lieu") + 1)
    assert jury.fill.fgColor.rgb.endswith("FDE68A")
    assert lieu.value == "HANKUY"
    assert not lieu.fill.fgColor.rgb.endswith("FDE68A")
